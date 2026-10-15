"""Import der Tabs REQ und DEL aus REQ_DEL_RAID_JIRA.xlsx als Epics ins Jira-Projekt SI.

Dry-Run (Standard): gibt fuer die ersten N Zeilen je Tab die Jira-Payload formatiert aus.
Es werden nur lesende Jira-Abfragen ausgefuehrt (Preflight), nichts wird geschrieben.

    py src/import_si_epics_from_excel.py                    # Dry-Run, 3 Zeilen je Tab
    py src/import_si_epics_from_excel.py --rows 10          # Dry-Run, 10 Zeilen je Tab
    py src/import_si_epics_from_excel.py --sheet DEL --rows 5

Ausfuehren (legt Epics an, transitioniert, setzt Reporter, erzeugt Links):

    py src/import_si_epics_from_excel.py --execute --rows 2   # Testlauf mit 2 Zeilen je Tab
    py src/import_si_epics_from_excel.py --execute            # alles

Mapping (Details siehe Konstanten unten):
    Initiative -> Parent | Title -> Summary | Description (+ Notes, Requirement/Deliverable,
    Workstream) -> Description | Priority -> Priority | Status -> Status (nach dem Anlegen
    transitioniert) | Assignee (erster Eintrag) -> Assignee | Reporter fix | Label fix |
    Component je Tab. Danach "relates to"-Links: DEL."Related Requirements" -> REQ."ID".

Wiederholbarkeit: angelegte Keys stehen in cache/si_epic_import_state.json; zusaetzlich
werden sie in eine Kopie der Excel-Datei (*_with_keys.xlsx) geschrieben. Bereits angelegte
Zeilen werden uebersprungen.

Credentials: Umgebungsvariablen ATLASSIAN_USER und ATLASSIAN_TOKEN.
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import openpyxl
import requests
from rich import box
from rich.console import Console
from rich.panel import Panel
from rich.rule import Rule
from rich.table import Table
from rich.text import Text

JIRA_BASE_URL = "https://trustedshops.atlassian.net"
PROJECT_KEY = "SI"
ISSUE_TYPE = "Epic"
INITIATIVE_TYPE = "Initiative"
LABEL = "Internal_AI_Environment"
REPORTER_NAME = "Stephan Marg"
LINK_TYPE = "Relates"

DEFAULT_XLSX = Path("tmp/REQ_DEL_RAID_JIRA.xlsx")
DEFAULT_STATE = Path("cache/si_epic_import_state.json")
DEFAULT_DRY_RUN_ROWS = 3

# Excel-Status (casefold) -> Status im SI-Workflow. Neue Epics starten in "Backlog".
STATUS_MAP = {"open": "ToDo", "in progress": "In Progress"}
# Excel-Prioritaet (casefold) -> Jira-Prioritaet.
PRIORITY_MAP = {"high": "High", "medium": "Medium", "low": "Low"}

EMPTY_MARKERS = {"", "-", "–", "—", "none"}


@dataclass(frozen=True)
class SheetSpec:
    name: str
    component: str
    id_col: str
    name_col: str  # Spalte, deren Wert als "<name_label>: <Wert>" in die Description kommt
    name_label: str
    workstream_col: str
    key_col: str  # Spalte, in die der Jira-Key zurueckgeschrieben wird
    color: str
    notes_col: str | None = None
    related_col: str | None = None  # Spalte mit Verweisen auf REQ-IDs (nur DEL)


SHEETS: dict[str, SheetSpec] = {
    "REQ": SheetSpec(
        name="REQ",
        component="Project Requirement",
        id_col="ID",
        name_col="Requirement",
        name_label="Requirement",
        workstream_col="Impacted Projects/Workstreams",
        key_col="JIRA",
        color="cyan",
        notes_col="Notes",
    ),
    "DEL": SheetSpec(
        name="DEL",
        component="Project Deliverable",
        id_col="DEL ID",
        name_col="Deliverable",
        name_label="Deliverable",
        workstream_col="Workstream",
        key_col="Jira Key",
        color="magenta",
        related_col="Related Requirements",
    ),
}


@dataclass
class EpicRow:
    spec: SheetSpec
    excel_row: int
    item_id: str
    initiative: str
    summary: str
    description: str
    notes: str
    name_value: str
    workstream: str
    priority: str
    status: str
    assignee_raw: str
    related_ids: list[str]


class JiraError(RuntimeError):
    def __init__(self, status: int, message: str):
        super().__init__(f"HTTP {status}: {message}")
        self.status = status
        self.message = message


# ---------------------------------------------------------------------------
# Excel einlesen (rein, ohne Jira)
# ---------------------------------------------------------------------------

def clean(value: Any) -> str:
    text = "" if value is None else str(value).strip()
    return "" if text.casefold() in EMPTY_MARKERS else text


def split_ids(value: str) -> list[str]:
    return [part.strip() for part in re.split(r"[;,]", value) if part.strip()]


def first_person(assignee_raw: str) -> str:
    """'A / B / C' -> 'A'. Mehrfach-Owner: nur der erste Eintrag zaehlt."""
    return assignee_raw.split("/")[0].strip()


def read_rows(xlsx: Path, spec: SheetSpec) -> list[EpicRow]:
    workbook = openpyxl.load_workbook(xlsx, data_only=True, read_only=True)
    try:
        rows = list(workbook[spec.name].iter_rows(values_only=True))
    finally:
        workbook.close()
    header = [clean(h) for h in rows[0]]

    def col(name: str) -> int:
        if name not in header:
            raise SystemExit(f"Tab {spec.name}: Spalte '{name}' nicht gefunden. Vorhanden: {header}")
        return header.index(name)

    idx = {n: col(n) for n in ("Initiative", spec.id_col, spec.name_col, "Title", "Description",
                               "Priority", "Status", spec.workstream_col, "Assignee")}
    notes_idx = col(spec.notes_col) if spec.notes_col else None
    related_idx = col(spec.related_col) if spec.related_col else None

    result: list[EpicRow] = []
    for number, row in enumerate(rows[1:], start=2):
        item_id = clean(row[idx[spec.id_col]])
        if not item_id:
            continue
        name_value = clean(row[idx[spec.name_col]])
        summary = clean(row[idx["Title"]]) or f"{item_id} {name_value}".strip()
        result.append(
            EpicRow(
                spec=spec,
                excel_row=number,
                item_id=item_id,
                initiative=clean(row[idx["Initiative"]]),
                summary=summary,
                description=clean(row[idx["Description"]]),
                notes=clean(row[notes_idx]) if notes_idx is not None else "",
                name_value=name_value,
                workstream=clean(row[idx[spec.workstream_col]]),
                priority=clean(row[idx["Priority"]]),
                status=clean(row[idx["Status"]]),
                assignee_raw=clean(row[idx["Assignee"]]),
                related_ids=split_ids(clean(row[related_idx])) if related_idx is not None else [],
            )
        )
    return result


# ---------------------------------------------------------------------------
# Payload bauen (rein)
# ---------------------------------------------------------------------------

def _paragraph(text: str, label: str | None = None) -> dict[str, Any]:
    content: list[dict[str, Any]] = []
    if label:
        content.append({"type": "text", "text": f"{label}: ", "marks": [{"type": "strong"}]})
    content.append({"type": "text", "text": text})
    return {"type": "paragraph", "content": content}


def build_description(row: EpicRow) -> dict[str, Any]:
    blocks: list[dict[str, Any]] = []
    if row.description:
        blocks.append(_paragraph(row.description))
    if row.notes:
        blocks.append(_paragraph(row.notes, "Notes"))
    if row.name_value:
        blocks.append(_paragraph(row.name_value, row.spec.name_label))
    if row.workstream:
        blocks.append(_paragraph(row.workstream, "Workstream"))
    return {"version": 1, "type": "doc", "content": blocks}


def target_status(row: EpicRow) -> str:
    try:
        return STATUS_MAP[row.status.casefold()]
    except KeyError:
        raise ValueError(f"{row.item_id}: kein Status-Mapping fuer '{row.status}'") from None


def jira_priority(row: EpicRow) -> str:
    try:
        return PRIORITY_MAP[row.priority.casefold()]
    except KeyError:
        raise ValueError(f"{row.item_id}: kein Priority-Mapping fuer '{row.priority}'") from None


def build_fields(row: EpicRow, *, reporter_id: str, assignee_id: str | None) -> dict[str, Any]:
    fields: dict[str, Any] = {
        "project": {"key": PROJECT_KEY},
        "issuetype": {"name": ISSUE_TYPE},
        "summary": row.summary,
        "description": build_description(row),
        "priority": {"name": jira_priority(row)},
        "components": [{"name": row.spec.component}],
        "labels": [LABEL],
        "reporter": {"accountId": reporter_id},
    }
    if row.initiative:
        fields["parent"] = {"key": row.initiative}
    if assignee_id:
        fields["assignee"] = {"accountId": assignee_id}
    return fields


def build_link_payload(del_key: str, req_key: str) -> dict[str, Any]:
    return {
        "type": {"name": LINK_TYPE},
        "inwardIssue": {"key": req_key},
        "outwardIssue": {"key": del_key},
    }


# ---------------------------------------------------------------------------
# Jira-Zugriff (duenn, bewusst ohne Pseudonymisierung: Account-IDs werden nie ausgegeben)
# ---------------------------------------------------------------------------

class Jira:
    def __init__(self) -> None:
        user, token = os.getenv("ATLASSIAN_USER"), os.getenv("ATLASSIAN_TOKEN")
        if not user or not token:
            raise SystemExit("ATLASSIAN_USER und ATLASSIAN_TOKEN muessen gesetzt sein.")
        self._session = requests.Session()
        self._session.auth = (user, token)
        self._session.headers.update({"Accept": "application/json"})

    def _request(self, method: str, path: str, **kwargs: Any) -> Any:
        for attempt in range(5):
            response = self._session.request(method, f"{JIRA_BASE_URL}{path}", **kwargs)
            if response.status_code == 429:
                time.sleep(float(response.headers.get("Retry-After", 2 ** attempt)))
                continue
            if not response.ok:
                raise JiraError(response.status_code, response.text[:400])
            return response.json() if response.content else None
        raise JiraError(429, "Rate limit nach 5 Versuchen nicht aufgeloest")

    def search(self, jql: str, fields: list[str]) -> list[dict[str, Any]]:
        issues: list[dict[str, Any]] = []
        token: str | None = None
        while True:
            payload: dict[str, Any] = {"jql": jql, "fields": fields, "maxResults": 100}
            if token:
                payload["nextPageToken"] = token
            data = self._request("POST", "/rest/api/3/search/jql", json=payload) or {}
            issues.extend(data.get("issues") or [])
            token = data.get("nextPageToken")
            if not token or not data.get("issues"):
                return issues

    def find_user(self, display_name: str, *, assignable: bool) -> str | None:
        """accountId bei genau einem exakten Treffer (displayName, aktiv, Atlassian-Konto)."""
        if assignable:
            path, params = "/rest/api/3/user/assignable/search", {"project": PROJECT_KEY, "query": display_name}
        else:
            path, params = "/rest/api/3/user/search", {"query": display_name}
        users = self._request("GET", path, params=params) or []
        exact = [
            u for u in users
            if u.get("active") and u.get("accountType") == "atlassian"
            and u.get("displayName", "").casefold() == display_name.casefold()
        ]
        return exact[0]["accountId"] if len(exact) == 1 else None

    def component_names(self) -> set[str]:
        return {c["name"] for c in self._request("GET", f"/rest/api/3/project/{PROJECT_KEY}/components")}

    def epic_status_names(self) -> set[str]:
        for entry in self._request("GET", f"/rest/api/3/project/{PROJECT_KEY}/statuses"):
            if entry["name"] == ISSUE_TYPE:
                return {s["name"] for s in entry["statuses"]}
        return set()

    def priority_names(self) -> set[str]:
        return {p["name"] for p in self._request("GET", "/rest/api/3/priority")}

    def create_issue(self, fields: dict[str, Any]) -> str:
        return self._request("POST", "/rest/api/3/issue", json={"fields": fields})["key"]

    def set_reporter(self, key: str, account_id: str) -> None:
        self._request("PUT", f"/rest/api/3/issue/{key}", json={"fields": {"reporter": {"accountId": account_id}}})

    def transition(self, key: str, to_status: str) -> None:
        data = self._request("GET", f"/rest/api/3/issue/{key}/transitions") or {}
        match = next(
            (t for t in data.get("transitions", []) if t["to"]["name"].casefold() == to_status.casefold()), None
        )
        if not match:
            available = [t["to"]["name"] for t in data.get("transitions", [])]
            raise JiraError(400, f"Kein Uebergang nach '{to_status}'. Verfuegbar: {available}")
        self._request("POST", f"/rest/api/3/issue/{key}/transitions", json={"transition": {"id": match["id"]}})

    def linked_keys(self, key: str) -> set[str]:
        issue = self._request("GET", f"/rest/api/3/issue/{key}", params={"fields": "issuelinks"})
        linked: set[str] = set()
        for link in issue["fields"].get("issuelinks", []):
            other = link.get("inwardIssue") or link.get("outwardIssue")
            if other:
                linked.add(other["key"])
        return linked

    def create_link(self, payload: dict[str, Any]) -> None:
        self._request("POST", "/rest/api/3/issueLink", json=payload)


# ---------------------------------------------------------------------------
# State (Wiederholbarkeit)
# ---------------------------------------------------------------------------

def load_state(path: Path) -> dict[str, Any]:
    if path.is_file():
        return json.loads(path.read_text(encoding="utf-8"))
    return {"epics": {}, "links": []}


def save_state(path: Path, state: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(state, indent=2, ensure_ascii=False), encoding="utf-8")


def write_keys_to_excel_copy(xlsx: Path, state: dict[str, Any]) -> Path | None:
    keys = {item_id: entry["key"] for item_id, entry in state["epics"].items()}
    if not keys:
        return None
    workbook = openpyxl.load_workbook(xlsx)  # Formeln bleiben erhalten (Werte rechnet Excel neu)
    for spec in SHEETS.values():
        sheet = workbook[spec.name]
        header = [clean(c.value) for c in sheet[1]]
        id_idx, key_idx = header.index(spec.id_col) + 1, header.index(spec.key_col) + 1
        for r in range(2, sheet.max_row + 1):
            key = keys.get(clean(sheet.cell(r, id_idx).value))
            if key:
                sheet.cell(r, key_idx).value = key
    target = xlsx.with_name(f"{xlsx.stem}_with_keys{xlsx.suffix}")
    workbook.save(target)
    return target


# ---------------------------------------------------------------------------
# Preflight (nur lesend)
# ---------------------------------------------------------------------------

@dataclass
class Preflight:
    problems: list[str]
    notes: list[str]
    reporter_id: str | None
    assignee_ids: dict[str, str | None]
    existing_by_summary: dict[str, str]


def run_preflight(jira: Jira, selected: list[EpicRow]) -> Preflight:
    problems: list[str] = []
    notes: list[str] = []

    # Initiatives
    wanted = sorted({r.initiative for r in selected if r.initiative})
    if wanted:
        found = {
            i["key"]: i["fields"]["issuetype"]["name"]
            for i in jira.search(f"key in ({','.join(wanted)})", ["issuetype"])
        }
        for key in wanted:
            if key not in found:
                problems.append(f"Parent {key} existiert nicht (oder keine Berechtigung).")
            elif found[key] != INITIATIVE_TYPE:
                problems.append(f"Parent {key} ist vom Typ '{found[key]}', erwartet '{INITIATIVE_TYPE}'.")
        notes.append(f"Parents geprueft: {', '.join(wanted)}")

    # Components, Status, Priority
    components = jira.component_names()
    for comp in sorted({r.spec.component for r in selected}):
        if comp not in components:
            problems.append(f"Component '{comp}' existiert nicht im Projekt {PROJECT_KEY}.")
    statuses = jira.epic_status_names()
    priorities = jira.priority_names()
    for row in selected:
        try:
            wanted_status, wanted_prio = target_status(row), jira_priority(row)
        except ValueError as exc:
            problems.append(str(exc))
            continue
        if wanted_status not in statuses:
            problems.append(f"{row.item_id}: Status '{wanted_status}' ist fuer {ISSUE_TYPE} in {PROJECT_KEY} nicht verfuegbar.")
        if wanted_prio not in priorities:
            problems.append(f"{row.item_id}: Prioritaet '{wanted_prio}' existiert nicht.")

    # User
    reporter_id = jira.find_user(REPORTER_NAME, assignable=False)
    if not reporter_id:
        problems.append(f"Reporter '{REPORTER_NAME}' nicht eindeutig gefunden.")
    assignee_ids: dict[str, str | None] = {}
    for row in selected:
        name = first_person(row.assignee_raw)
        if name and name not in assignee_ids:
            assignee_ids[name] = jira.find_user(name, assignable=True)

    # Bereits vorhandene Epics mit dem Import-Label (Schutz, falls die State-Datei fehlt)
    existing = {
        i["fields"]["summary"]: i["key"]
        for i in jira.search(f'project = {PROJECT_KEY} AND issuetype = {ISSUE_TYPE} AND labels = "{LABEL}"', ["summary"])
    }
    return Preflight(problems, notes, reporter_id, assignee_ids, existing)


# ---------------------------------------------------------------------------
# Ausgabe
# ---------------------------------------------------------------------------

def render_preflight(console: Console, pre: Preflight, selected: list[EpicRow]) -> None:
    table = Table(box=box.SIMPLE_HEAD, show_edge=False, pad_edge=False)
    table.add_column("Pruefung", style="bold")
    table.add_column("Ergebnis")
    for note in pre.notes:
        table.add_row("OK", Text(note, style="green"))
    table.add_row(
        "Reporter",
        Text(f"{REPORTER_NAME}: gefunden", style="green") if pre.reporter_id
        else Text(f"{REPORTER_NAME}: NICHT gefunden", style="bold red"),
    )
    missing = sorted(n for n, i in pre.assignee_ids.items() if not i)
    table.add_row(
        "Assignees",
        Text(f"{len(pre.assignee_ids) - len(missing)} gefunden", style="green")
        + (Text(f", ohne Treffer (Assignee bleibt leer): {', '.join(missing)}", style="yellow") if missing else Text("")),
    )
    clashes = [r.item_id for r in selected if r.summary in pre.existing_by_summary]
    table.add_row(
        "Duplikate",
        Text(f"{len(clashes)} Summary-Treffer unter vorhandenen Epics mit Label {LABEL}: {', '.join(clashes)}", style="yellow")
        if clashes else Text(f"keine Summary-Treffer unter vorhandenen Epics mit Label {LABEL}", style="green"),
    )
    for problem in pre.problems:
        table.add_row("FEHLER", Text(problem, style="bold red"))
    border = "red" if pre.problems else "green"
    console.print(Panel(table, title="[bold]Preflight (nur lesend)[/bold]", border_style=border, box=box.ROUNDED))


def adf_to_text(doc: dict[str, Any]) -> Text:
    """ADF-Description als lesbarer Text (fette Labels bleiben fett)."""
    result = Text()
    for index, paragraph in enumerate(doc["content"]):
        if index:
            result.append("\n")
        for node in paragraph["content"]:
            bold = any(m["type"] == "strong" for m in node.get("marks", []))
            result.append(node["text"], style="bold" if bold else "")
    return result


def render_row_panel(console: Console, row: EpicRow, pre: Preflight, state: dict[str, Any],
                     req_by_id: dict[str, EpicRow]) -> None:
    spec = row.spec
    person = first_person(row.assignee_raw)
    assignee_found = bool(person and pre.assignee_ids.get(person))

    fields = Table.grid(padding=(0, 2))
    fields.add_column(style=f"bold {spec.color}", no_wrap=True)
    fields.add_column()
    fields.add_row("Projekt / Typ", f"{PROJECT_KEY} / {ISSUE_TYPE}")
    fields.add_row("Summary", Text(row.summary, style="bold"))
    fields.add_row("Parent", row.initiative or "-")
    fields.add_row("Component", spec.component)
    fields.add_row("Priority", f"{jira_priority(row)}  [dim](Excel: {row.priority})[/dim]")
    fields.add_row("Status", f"Backlog -> [bold]{target_status(row)}[/bold]  [dim](Excel: {row.status})[/dim]")
    fields.add_row("Label", LABEL)
    fields.add_row("Reporter", REPORTER_NAME)
    if not person:
        fields.add_row("Assignee", "[dim]leer[/dim]")
    elif assignee_found:
        fields.add_row("Assignee", person + (f"  [dim](Excel: {row.assignee_raw})[/dim]" if row.assignee_raw != person else ""))
    else:
        fields.add_row("Assignee", f"[yellow]leer - kein User '{person}' gefunden[/yellow]  [dim](Excel: {row.assignee_raw})[/dim]")
    fields.add_row("Description", adf_to_text(build_description(row)))
    if row.related_ids:
        links = Text()
        for index, req_id in enumerate(row.related_ids):
            if index:
                links.append("\n")
            req = req_by_id.get(req_id)
            links.append("relates to ", style="dim")
            if req:
                links.append(req.summary)
            else:
                links.append(f"{req_id} (nicht im Tab REQ!)", style="bold red")
        fields.add_row("Links", links)
    existing = state["epics"].get(row.item_id, {}).get("key") or pre.existing_by_summary.get(row.summary)
    if existing:
        fields.add_row("Hinweis", f"[bold yellow]existiert bereits als {existing} -> wuerde uebersprungen[/bold yellow]")

    console.print(
        Panel(
            fields,
            title=f"[bold {spec.color}]{spec.name}[/bold {spec.color}] · Zeile {row.excel_row} · [bold]{row.item_id}[/bold]",
            border_style=spec.color,
            box=box.HEAVY,
            padding=(1, 2),
        )
    )


# ---------------------------------------------------------------------------
# Ablauf
# ---------------------------------------------------------------------------

def select_rows(all_rows: dict[str, list[EpicRow]], sheet: str, limit: int | None) -> dict[str, list[EpicRow]]:
    names = list(SHEETS) if sheet == "all" else [sheet]
    return {n: (all_rows[n][:limit] if limit is not None else all_rows[n]) for n in names}


def run_dry(console: Console, jira: Jira, all_rows: dict[str, list[EpicRow]], chosen: dict[str, list[EpicRow]],
            state: dict[str, Any]) -> int:
    selected = [r for rows in chosen.values() for r in rows]
    console.print(Rule("[bold]DRY-RUN[/bold] - es wird nichts in Jira geschrieben", style="yellow"))
    pre = run_preflight(jira, selected)
    render_preflight(console, pre, selected)
    req_by_id = {r.item_id: r for r in all_rows["REQ"]}
    for name, rows in chosen.items():
        console.print(Rule(f"[bold]{name}[/bold]: {len(rows)} von {len(all_rows[name])} Zeilen", style=SHEETS[name].color))
        for row in rows:
            render_row_panel(console, row, pre, state, req_by_id)
    console.print(Rule(style="yellow"))
    console.print("[bold yellow]Dry-Run beendet.[/bold yellow] Mit [bold]--execute[/bold] tatsaechlich anlegen.")
    return 1 if pre.problems else 0


def run_execute(console: Console, jira: Jira, xlsx: Path, state_path: Path, all_rows: dict[str, list[EpicRow]],
                chosen: dict[str, list[EpicRow]], state: dict[str, Any], assume_yes: bool) -> int:
    selected = [r for rows in chosen.values() for r in rows]
    pre = run_preflight(jira, selected)
    render_preflight(console, pre, selected)
    if pre.problems or not pre.reporter_id:
        console.print("[bold red]Abbruch: Preflight meldet Fehler.[/bold red]")
        return 1

    todo = [r for r in selected if r.item_id not in state["epics"]]
    console.print(
        f"{len(selected)} Zeilen gewaehlt, davon {len(selected) - len(todo)} schon angelegt (State). "
        f"Es werden bis zu [bold]{len(todo)}[/bold] Epics in {PROJECT_KEY} angelegt."
    )
    if not assume_yes and input("Fortfahren? [j/N] ").strip().casefold() not in {"j", "ja", "y", "yes"}:
        console.print("Abgebrochen.")
        return 1

    errors: list[str] = []
    for row in selected:
        entry = state["epics"].get(row.item_id)
        try:
            if entry is None:
                adopted = pre.existing_by_summary.get(row.summary)
                if adopted:
                    state["epics"][row.item_id] = {"key": adopted, "sheet": row.spec.name, "summary": row.summary,
                                                   "status_set": True, "reporter_set": True, "adopted": True}
                    save_state(state_path, state)
                    console.print(f"[yellow]=[/yellow] {row.item_id} existiert bereits als {adopted} (Summary-Treffer), uebernommen")
                    continue
                person = first_person(row.assignee_raw)
                assignee_id = pre.assignee_ids.get(person) if person else None
                fields = build_fields(row, reporter_id=pre.reporter_id, assignee_id=assignee_id)
                reporter_in_create = True
                try:
                    key = jira.create_issue(fields)
                except JiraError as exc:
                    if exc.status == 400 and "reporter" in exc.message.casefold():
                        fields.pop("reporter")
                        reporter_in_create = False
                        key = jira.create_issue(fields)
                    else:
                        raise
                entry = {"key": key, "sheet": row.spec.name, "summary": row.summary,
                         "status_set": False, "reporter_set": reporter_in_create}
                state["epics"][row.item_id] = entry
                save_state(state_path, state)
                console.print(f"[green]+[/green] {row.item_id} -> [bold]{key}[/bold]  {row.summary}")
            # Status und Reporter sind unabhaengige Nachbearbeitungen: ein Fehler stoppt nicht den anderen.
            if not entry["status_set"]:
                try:
                    jira.transition(entry["key"], target_status(row))
                    entry["status_set"] = True
                except (JiraError, ValueError) as exc:
                    errors.append(f"{row.item_id} ({entry['key']}): Status: {exc}")
                    console.print(f"[bold red]![/bold red] {row.item_id}: Status: {exc}")
            if not entry["reporter_set"]:
                try:
                    jira.set_reporter(entry["key"], pre.reporter_id)
                    entry["reporter_set"] = True
                except JiraError as exc:
                    errors.append(f"{row.item_id} ({entry['key']}): Reporter: {exc}")
                    console.print(f"[bold red]![/bold red] {row.item_id}: Reporter: {exc}")
            save_state(state_path, state)
        except (JiraError, ValueError) as exc:
            errors.append(f"{row.item_id}: {exc}")
            console.print(f"[bold red]![/bold red] {row.item_id}: {exc}")
            save_state(state_path, state)

    # Links: DEL.Related Requirements -> REQ.ID (nur fuer ausgewaehlte DEL-Zeilen)
    link_count = 0
    for row in chosen.get("DEL", []):
        del_entry = state["epics"].get(row.item_id)
        if not del_entry or not row.related_ids:
            continue
        try:
            already = jira.linked_keys(del_entry["key"])
        except JiraError as exc:
            errors.append(f"{row.item_id}: Links nicht lesbar: {exc}")
            continue
        for req_id in row.related_ids:
            req_entry = state["epics"].get(req_id)
            pair = f"{row.item_id}|{req_id}"
            if not req_entry:
                errors.append(f"{row.item_id}: REQ {req_id} nicht angelegt/ausgewaehlt, Link uebersprungen")
                continue
            if pair in state["links"] or req_entry["key"] in already:
                if pair not in state["links"]:
                    state["links"].append(pair)
                continue
            try:
                jira.create_link(build_link_payload(del_entry["key"], req_entry["key"]))
                state["links"].append(pair)
                save_state(state_path, state)
                link_count += 1
                console.print(f"[blue]~[/blue] {del_entry['key']} relates to {req_entry['key']}  ({pair})")
            except JiraError as exc:
                errors.append(f"{pair}: {exc}")
                console.print(f"[bold red]![/bold red] Link {pair}: {exc}")

    copy = write_keys_to_excel_copy(xlsx, state)
    console.print(Rule("[bold]Fertig[/bold]"))
    console.print(f"Epics im State: {len(state['epics'])}, neue Links: {link_count}, Fehler: {len(errors)}")
    if copy:
        console.print(f"Keys in Excel-Kopie: {copy}")
    for message in errors:
        console.print(f"[red]- {message}[/red]")
    return 1 if errors else 0


def main() -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description="Import REQ/DEL-Tabs aus Excel als Epics ins Jira-Projekt SI")
    parser.add_argument("--xlsx", type=Path, default=DEFAULT_XLSX)
    parser.add_argument("--state-file", type=Path, default=DEFAULT_STATE)
    parser.add_argument("--sheet", choices=["all", *SHEETS], default="all", help="Tab waehlen (Standard: beide)")
    parser.add_argument("--rows", type=int, default=None,
                        help=f"Anzahl Zeilen je Tab (Dry-Run Standard: {DEFAULT_DRY_RUN_ROWS}, --execute Standard: alle)")
    parser.add_argument("--execute", action="store_true", help="Tatsaechlich in Jira schreiben (ohne Flag: Dry-Run)")
    parser.add_argument("--yes", action="store_true", help="Rueckfrage vor --execute ueberspringen")
    args = parser.parse_args()

    console = Console(highlight=False)
    all_rows = {name: read_rows(args.xlsx, spec) for name, spec in SHEETS.items()}
    limit = args.rows if args.rows is not None else (None if args.execute else DEFAULT_DRY_RUN_ROWS)
    chosen = select_rows(all_rows, args.sheet, limit)
    state = load_state(args.state_file)
    jira = Jira()

    if args.execute:
        return run_execute(console, jira, args.xlsx, args.state_file, all_rows, chosen, state, args.yes)
    return run_dry(console, jira, all_rows, chosen, state)


if __name__ == "__main__":
    sys.exit(main())
